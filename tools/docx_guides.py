"""Editable editions using Google Docs-friendly fonts and native Word structure."""
from pathlib import Path
from tempfile import TemporaryDirectory
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from google_docs_title_sanitize import sanitize_docx, check_docx
from pdf_guides import link_target, slug


def bookmark(paragraph, title, identifier):
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(identifier))
    start.set(qn('w:name'), slug(title).replace('-', '_'))
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(identifier))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_link(paragraph, label, target, bold=False, italic=False):
    hyperlink = OxmlElement('w:hyperlink')
    if target.startswith('#'):
        hyperlink.set(qn('w:anchor'), target[1:].replace('-', '_'))
    else:
        hyperlink.set(qn('r:id'), paragraph.part.relate_to(link_target(target), RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '2F4F3A')
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    properties.extend([color, underline])
    for enabled, tag in [(bold, 'w:b'), (italic, 'w:i')]:
        if enabled: properties.append(OxmlElement(tag))
    run.append(properties)
    text = OxmlElement('w:t')
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def inline(paragraph, text, bold=False, italic=False):
    """Retain emphasis and links instead of flattening Markdown into plain text."""
    tokens = re.compile(r'\*\*(.+?)\*\*|\*([^*]+)\*|\[([^\]]+)\]\(([^)]+)\)')
    position = 0
    for token in tokens.finditer(text):
        run = paragraph.add_run(text[position:token.start()])
        run.bold, run.italic = bold, italic
        if token[1] is not None:
            inline(paragraph, token[1], True, italic)
        elif token[2] is not None:
            inline(paragraph, token[2], bold, True)
        else:
            add_link(paragraph, token[3], token[4], bold, italic)
        position = token.end()
    run = paragraph.add_run(text[position:])
    run.bold, run.italic = bold, italic


def number_list(document):
    numbering = document.part.numbering_part.element
    abstract_id = max([int(element.get(qn('w:abstractNumId'))) for element in numbering.findall(qn('w:abstractNum'))], default=0) + 1
    number_id = max([int(element.get(qn('w:numId'))) for element in numbering.findall(qn('w:num'))], default=0) + 1
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), str(abstract_id))
    level = OxmlElement('w:lvl')
    level.set(qn('w:ilvl'), '0')
    for tag, value in [('w:start','1'), ('w:numFmt','decimal'), ('w:lvlText','%1.'), ('w:lvlJc','left')]:
        element = OxmlElement(tag)
        element.set(qn('w:val'), value)
        level.append(element)
    properties = OxmlElement('w:pPr')
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'),'num')
    tab.set(qn('w:pos'),'360')
    tabs.append(tab)
    properties.append(tabs)
    indent = OxmlElement('w:ind')
    indent.set(qn('w:left'),'360')
    indent.set(qn('w:hanging'),'360')
    properties.append(indent)
    level.append(properties)
    abstract.append(level)
    numbering.append(abstract)
    number = OxmlElement('w:num')
    number.set(qn('w:numId'), str(number_id))
    reference = OxmlElement('w:abstractNumId')
    reference.set(qn('w:val'), str(abstract_id))
    number.append(reference)
    numbering.append(number)
    return number_id


def numbered_paragraph(document, text, number_id):
    paragraph = document.add_paragraph(style='List Number')
    properties = paragraph._p.get_or_add_pPr()
    number = OxmlElement('w:numPr')
    for tag, value in [('w:ilvl','0'), ('w:numId',str(number_id))]:
        element = OxmlElement(tag)
        element.set(qn('w:val'), value)
        number.append(element)
    properties.append(number)
    inline(paragraph, text)


def configure(document, label):
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    section.header_distance = section.footer_distance = Inches(.3)
    for name in ['Normal','List Bullet','List Number','Caption','Title','Subtitle','Heading 1','Heading 2']:
        style = document.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor.from_string('000000')
        style.font.size = Pt(11)
        style.paragraph_format.line_spacing = 1.12
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.widow_control = True
        for font_properties in style.element.iter(qn('w:rFonts')):
            for attribute in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                font_properties.attrib.pop(qn('w:'+attribute),None)
    for name,size in [('Title',30), ('Heading 1',22), ('Heading 2',15)]:
        style = document.styles[name]
        style.font.name = 'Georgia'
        style.font.size = Pt(size)
        style.font.bold = name != 'Title'
        style.paragraph_format.space_before = Pt(18 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(10)
        style.paragraph_format.keep_with_next = True
    document.styles['Caption'].font.size = Pt(9)
    document.styles['Caption'].font.color.rgb = RGBColor.from_string('6B6B60')
    for name in ['List Bullet','List Number']:
        document.styles[name].paragraph_format.space_after = Pt(5)
        document.styles[name].paragraph_format.left_indent = Inches(.25)
        document.styles[name].paragraph_format.first_line_indent = Inches(-.25)
    header = section.header.paragraphs[0]
    header.text = f'TRAVEL MONKEY  /  {label.upper()}'
    header.style = document.styles['Normal']
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string('000000')
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run('Travel Monkey  /  ')
    run.font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    document.core_properties.author = 'Travel Monkey'
    document.core_properties.title = f'Travel Monkey {label.lower()}'


def table(document, rows):
    grid = document.add_table(rows=0, cols=2)
    grid.alignment = WD_TABLE_ALIGNMENT.LEFT
    grid.autofit = False
    widths = [Inches(1.7), Inches(5.3)]
    for column,width in zip(grid.columns,widths): column.width = width
    properties = grid._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        edge = OxmlElement('w:'+side)
        for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]: edge.set(qn('w:'+key),value)
        borders.append(edge)
    properties.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for side in ['top','left','bottom','right']:
        margin = OxmlElement('w:'+side)
        margin.set(qn('w:w'),'100')
        margin.set(qn('w:type'),'dxa')
        margins.append(margin)
    properties.append(margins)
    for row_number,values in enumerate(rows):
        row = grid.add_row()
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement('w:cantSplit'))
        if row_number == 0: row_properties.append(OxmlElement('w:tblHeader'))
        for cell,value,width in zip(row.cells,values,widths):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.06
            paragraph.paragraph_format.keep_with_next = row_number==0 or (len(rows)<=6 and row_number<len(rows)-1)
            inline(paragraph,value,bold=row_number==0)
            for run in paragraph.runs: run.font.size = Pt(10)
            if row_number == 0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'),'EDEDED')
                cell._tc.get_or_add_tcPr().append(shading)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.add_run().font.size = Pt(3)


def add_image(document, path, alt, compact=False):
    with Image.open(path) as picture:
        width,height = picture.size
    scale = min(5.9/width, (2.6 if compact else 3.5)/height)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_after = Pt(4)
    shape = paragraph.add_run().add_picture(str(path),width=Inches(width*scale),height=Inches(height*scale))
    shape._inline.docPr.set('descr',alt)


def build_docx(source_root, output_directory, kind):
    document = Document()
    label = 'Quick start' if kind=='quick-start' else 'User guide'
    configure(document,label)
    logo = document.add_paragraph()
    logo.add_run().add_picture(str(source_root/'tools/assets/monkey-mark.png'),width=Inches(.5))
    document.add_paragraph('Travel Monkey '+label.lower(),style='Title')
    markdown = (source_root/f'{kind}.md').read_text(encoding='utf-8')
    lines = markdown.splitlines()[1:]
    index, identifier, number_id = 0, 0, None
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            number_id = None
            index += 1
            continue
        if line.startswith('## '):
            title = line[3:]
            if title in ['Contents','Before you begin'] or (kind=='quick-start' and title.startswith(('2.','4.'))):
                document.add_page_break()
            paragraph = document.add_paragraph(title,style='Heading 1')
            bookmark(paragraph,title,identifier)
            identifier += 1
            index += 1
            continue
        if line.startswith('### '):
            paragraph = document.add_paragraph(line[4:],style='Heading 2')
            bookmark(paragraph,line[4:],identifier)
            identifier += 1
            index += 1
            continue
        picture = re.match(r'!\[([^\]]*)\]\(([^)]+)\)',line)
        if picture:
            add_image(document,source_root/picture[2],picture[1],compact=kind=='quick-start')
            index += 1
            continue
        if line.startswith('|'):
            rows=[]
            while index<len(lines) and lines[index].strip().startswith('|'):
                values = [value.strip() for value in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',value) for value in values): rows.append(values)
                index += 1
            table(document,rows)
            continue
        item = re.match(r'(- |\d+\. )(.*)',line)
        content = [item[2] if item else line]
        index += 1
        while index<len(lines) and lines[index].strip() and not re.match(r'(- |\d+\. |## |### |!\[|\|)',lines[index]):
            content.append(lines[index].strip())
            index += 1
        text = ' '.join(content)
        if item and not item[1].startswith('-'):
            if number_id is None: number_id = number_list(document)
            numbered_paragraph(document,text,number_id)
        else:
            style = 'List Bullet' if item else ('Caption' if text.startswith('*') and not text.startswith('**') else 'Normal')
            paragraph = document.add_paragraph(style=style)
            inline(paragraph,text)
            next_line = next((following.strip() for following in lines[index:] if following.strip()),'')
            if next_line.startswith('|'):
                paragraph.paragraph_format.keep_with_next = True
    output_directory.mkdir(parents=True,exist_ok=True)
    filename = output_directory/f'travel-monkey-{kind}.docx'
    with TemporaryDirectory(prefix='travel-monkey-docx-') as temporary:
        raw = Path(temporary)/'guide.docx'
        document.save(raw)
        sanitize_docx(raw,filename,leading_nonempty_paragraphs=3)
    issues = check_docx(filename,leading_nonempty_paragraphs=3)
    if issues: raise ValueError(issues)
    print(filename.name, 'editable document created and title audit passed')


def build_docxs(source_root, output_directory):
    for kind in ['quick-start','user-guide']:
        build_docx(source_root,output_directory,kind)
